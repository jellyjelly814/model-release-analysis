from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
src = ROOT / "GLM-5.3_研究报告.md"
out = ROOT / "GLM-5.3_竞品分析_20260820.docx"
lines = src.read_text(encoding="utf-8").splitlines()
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
styles["Normal"].font.size = Pt(10.5)
for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13)):
    styles[name].font.name = "Arial"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

def clean_inline(s):
    s = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", s)
    s = re.sub(r"\*\*([^*]+)\*\*", r"\1", s)
    s = re.sub(r"`([^`]+)`", r"\1", s)
    return s.strip()

def add_para(text, style=None, bold=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(clean_inline(text))
    r.font.color.rgb = RGBColor(0, 0, 0)
    r.bold = bold
    return p

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def add_table(rows):
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for j, val in enumerate(rows[0]):
        c = table.rows[0].cells[j]
        c.text = clean_inline(val)
        shade(c, "EAF2FF")
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in c.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(0, 0, 0)
    for row in rows[1:]:
        cells = table.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = clean_inline(val)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[j].paragraphs:
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph()

def add_image(path, caption):
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(path), width=Inches(6.4))
        cp = doc.add_paragraph(caption)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.runs[0].italic = True
        cp.runs[0].font.color.rgb = RGBColor(0, 0, 0)

i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip():
        i += 1
        continue
    if line.startswith("> "):
        add_para(line[2:], bold=True)
        i += 1
        continue
    if line.startswith("# "):
        add_para(line[2:], style="Title")
        i += 1
        continue
    if line.startswith("## "):
        add_para(line[3:], style="Heading 1")
        i += 1
        continue
    if line.startswith("### "):
        add_para(line[4:], style="Heading 2")
        if line.startswith("### 4.1"):
            add_image(ROOT / "images/S1UsYfhUGl.png", "图 1：Z.ai 官方工程与 Agent 基准对比")
        elif line.startswith("### 4.2"):
            add_image(ROOT / "images/SyHnYGhUMg.png", "图 2：Z.ai Code Bench 的推理强度与输出 Token")
        elif line.startswith("### 4.3"):
            add_image(ROOT / "images/20260814-124300.png", "图 3：Z.ai 官方网络安全评测")
        i += 1
        continue
    if line.startswith("- ") or re.match(r"\d+\. ", line):
        text = re.sub(r"^(?:- |\d+\. )", "", line)
        p = doc.add_paragraph(style="List Bullet" if line.startswith("-") else "List Number")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(clean_inline(text))
        r.font.color.rgb = RGBColor(0, 0, 0)
        i += 1
        continue
    if line.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            vals = [x.strip() for x in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r"[-: ]+", x or " ") for x in vals):
                rows.append(vals)
            i += 1
        add_table(rows)
        continue
    if line.startswith("!["):
        m = re.search(r"\(([^)]+)\)", line)
        if m:
            pth = ROOT / m.group(1)
            add_image(pth, clean_inline(line.split("]", 1)[0].lstrip("![")))
        i += 1
        continue
    if line.startswith("```"):
        i += 1
        code = []
        while i < len(lines) and not lines[i].startswith("```"):
            code.append(lines[i]); i += 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        r = p.add_run("\n".join(code))
        r.font.name = "Menlo"; r.font.size = Pt(8); r.font.color.rgb = RGBColor(0, 0, 0)
        i += 1
        continue
    add_para(line)
    i += 1

doc.core_properties.title = "GLM-5.3 竞品分析"
doc.core_properties.subject = "智谱 GLM-5.3 发布、能力、生态与风险研究"
doc.core_properties.author = "Codex"
doc.save(out)
print(out)
